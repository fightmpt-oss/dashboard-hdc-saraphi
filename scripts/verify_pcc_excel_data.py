import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('PCC 2569/PCC_69_R.1.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)
ws = wb['เขต 1']
print(f'Max rows: {ws.max_row}, Max cols: {ws.max_column}')

saraphi_codes = {
    '11135', '06014', '06015', '06016', '06017', '06018',
    '06020', '06021', '06022', '06023', '06024', '13994', '14461', '99758'
}

found_rows = []
for r in range(2, ws.max_row + 1):
    hsub = str(ws.cell(r, 9).value or '').strip()
    # Normalize 4 or 5 digit code with leading zero if needed
    if len(hsub) == 4 and ('0' + hsub) in saraphi_codes:
        hsub = '0' + hsub
    if hsub in saraphi_codes or (ws.cell(r, 8).value and 'สารภี' in str(ws.cell(r, 8).value)):
        found_rows.append((
            r,
            hsub,
            ws.cell(r, 10).value,
            ws.cell(r, 11).value,
            ws.cell(r, 12).value,
            ws.cell(r, 45).value
        ))

print(f'\nFound {len(found_rows)} rows matching Saraphi in Excel:')
total_excel_budget = 0.0
for fr in found_rows:
    b = float(fr[5] or 0)
    total_excel_budget += b
    print(f'  Row {fr[0]:4d}: HSUB={fr[1]} | Name={fr[2]} | UC_pop={fr[3]} | UC35_pop={fr[4]} | Total_budget={b:,.2f}')

print(f'\nTotal Excel Budget for Saraphi: {total_excel_budget:,.2f} บาท')

# Compare with data/pcc_2569_master.json
with open('data/pcc_2569_master.json', 'r', encoding='utf-8') as f:
    master = json.load(f)

dist = master.get('saraphi_district', {})
print('\nDistrict in master.json:')
print(f'  UC35_pop: {dist.get("uc35_pop")}')
print(f'  Total Budget: {dist.get("total_budget"):,.2f} บาท')
for k in ['kpi1', 'kpi2', 'kpi3', 'kpi4']:
    kdata = dist.get(k, {})
    print(f'  {k}: a={kdata.get("a"):,} | b={kdata.get("b"):,} | rate={kdata.get("rate")}% | budget={kdata.get("budget"):,.2f}')

units = master.get('units', {})
print(f'\nUnits in master.json count: {len(units)}')
for hc, u in units.items():
    print(f'  {hc}: {u.get("name")} | Total={u.get("total_budget"):,.2f} บ. | UC35={u.get("uc35_pop")} | k1={u["kpi1"]["rate"]}% | k2={u["kpi2"]["rate"]}% | k3={u["kpi3"]["rate"]}% | k4={u["kpi4"]["rate"]}%')
